import os
import re
from datetime import datetime
from threading import Lock

from flask import (
    Flask,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    url_for,
)
from openpyxl import Workbook, load_workbook

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-secret-key")

DATA_DIR = os.environ.get("DATA_DIR", os.path.join(os.path.dirname(__file__), "data"))
EXCEL_PATH = os.path.join(DATA_DIR, "miembros_mesa.xlsx")
HEADERS = ["DNI", "Región", "Provincia", "Distrito", "Dirección del local de votación", "Fecha de registro"]

excel_lock = Lock()

REGIONES_PERU = [
    "Amazonas", "Áncash", "Apurímac", "Arequipa", "Ayacucho", "Cajamarca",
    "Callao", "Cusco", "Huancavelica", "Huánuco", "Ica", "Junín",
    "La Libertad", "Lambayeque", "Lima", "Loreto", "Madre de Dios",
    "Moquegua", "Pasco", "Piura", "Puno", "San Martín", "Tacna",
    "Tumbes", "Ucayali",
]

DNI_PATTERN = re.compile(r"^\d{8}$")


def init_excel():
    os.makedirs(DATA_DIR, exist_ok=True)
    if not os.path.exists(EXCEL_PATH):
        wb = Workbook()
        ws = wb.active
        ws.title = "Miembros de Mesa"
        ws.append(HEADERS)
        wb.save(EXCEL_PATH)


def dni_ya_registrado(dni):
    wb = load_workbook(EXCEL_PATH)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and str(row[0]) == dni:
            return True
    return False


def guardar_miembro(datos):
    with excel_lock:
        init_excel()
        wb = load_workbook(EXCEL_PATH)
        ws = wb.active
        ws.append(
            [
                datos["dni"],
                datos["region"],
                datos["provincia"],
                datos["distrito"],
                datos["direccion"],
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            ]
        )
        wb.save(EXCEL_PATH)


def leer_miembros():
    init_excel()
    wb = load_workbook(EXCEL_PATH)
    ws = wb.active
    return list(ws.iter_rows(min_row=2, values_only=True))


@app.route("/")
def index():
    return render_template("index.html", regiones=REGIONES_PERU)


@app.route("/registrar", methods=["POST"])
def registrar():
    dni = request.form.get("dni", "").strip()
    region = request.form.get("region", "").strip()
    provincia = request.form.get("provincia", "").strip()
    distrito = request.form.get("distrito", "").strip()
    direccion = request.form.get("direccion", "").strip()

    errores = []
    if not DNI_PATTERN.match(dni):
        errores.append("El DNI debe tener exactamente 8 dígitos numéricos.")
    if not region:
        errores.append("La región es obligatoria.")
    if not provincia:
        errores.append("La provincia es obligatoria.")
    if not distrito:
        errores.append("El distrito es obligatorio.")
    if not direccion:
        errores.append("La dirección del local de votación es obligatoria.")

    if not errores:
        with excel_lock:
            init_excel()
            if dni_ya_registrado(dni):
                errores.append(f"El DNI {dni} ya se encuentra registrado.")

    if errores:
        for error in errores:
            flash(error, "error")
        return redirect(url_for("index"))

    guardar_miembro(
        {
            "dni": dni,
            "region": region,
            "provincia": provincia,
            "distrito": distrito,
            "direccion": direccion,
        }
    )
    flash("Miembro de mesa registrado correctamente.", "success")
    return redirect(url_for("index"))


@app.route("/miembros")
def miembros():
    filas = leer_miembros()
    return render_template("miembros.html", miembros=filas)


@app.route("/descargar")
def descargar():
    init_excel()
    return send_file(EXCEL_PATH, as_attachment=True, download_name="miembros_mesa.xlsx")


@app.route("/healthz")
def healthz():
    return {"status": "ok"}, 200


if __name__ == "__main__":
    init_excel()
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=os.environ.get("FLASK_DEBUG", "0") == "1")
